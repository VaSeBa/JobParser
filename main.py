"""
Парсер вакансий с hh.ru с графическим интерфейсом
Автор: VaSeBa
Версия: 2.3 (исправленная)
"""

import os
import sys
import time
from datetime import datetime, timedelta
from typing import Generator, Tuple, Optional, Dict, Any

import requests
import pandas as pd
from dateutil import parser  # Для парсинга дат

from PySide6.QtWidgets import (
    QApplication, QWidget, QVBoxLayout, QLabel,
    QLineEdit, QPushButton, QTextEdit, QProgressBar
)
from PySide6.QtCore import QThread, Signal, Qt
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


class ParserThread(QThread):
    """Поток для выполнения парсинга без блокировки GUI"""
    update_log = Signal(str)  # Сигнал для обновления лога
    finished = Signal()  # Сигнал завершения работы
    progress_updated = Signal(int)  # Сигнал обновления прогресса

    def __init__(self, profession: str):
        super().__init__()
        self.profession = profession
        self.running = True  # Флаг для остановки потока
        self.total_days = 30  # Период поиска (дней)

    def run(self) -> None:
        """Основной метод выполнения потока"""
        try:
            end_date = datetime.now().replace(
                hour=0, minute=0, second=0, microsecond=0
            )
            start_date = end_date - timedelta(days=self.total_days)

            # Генерация интервалов дат для поиска
            date_intervals = list(self.date_range(start_date, end_date))
            total_intervals = len(date_intervals)

            all_vacancies = []

            for i, (date_from, date_to) in enumerate(date_intervals):
                if not self.running:
                    return  # Выход при остановке

                # Расчет и отправка прогресса
                progress = int((i + 1) / total_intervals * 100)
                self.progress_updated.emit(progress)

                # Парсинг вакансий за период
                self.update_log.emit(
                    f"Парсинг за период: {date_from.date()} - {date_to.date()}"
                )
                vacancies = self.get_vacancies(
                    text=self.profession,
                    date_from=date_from.isoformat(),
                    date_to=date_to.isoformat()
                )

                all_vacancies.extend(vacancies)
                self.update_log.emit(f"Найдено вакансий: {len(vacancies)}")
                time.sleep(0.5)  # Задержка для избежания блокировки

            # Сохранение результатов
            if all_vacancies:
                self.save_to_excel(all_vacancies)
                self.update_log.emit(f"\nВсего собрано вакансий: {len(all_vacancies)}")
            else:
                self.update_log.emit("\nНе найдено подходящих вакансий")

            self.progress_updated.emit(100)

        except Exception as e:
            self.update_log.emit(f"Критическая ошибка: {str(e)}")
        finally:
            self.finished.emit()

    def get_vacancies(self, text: str, date_from: str, date_to: str) -> list:
        """Получение вакансий через API hh.ru"""
        base_url = "https://api.hh.ru/vacancies"
        params = {
            "text": text,
            "area": 113,  # 113 = Россия
            "date_from": date_from,
            "date_to": date_to,
            "per_page": 100,  # Максимальное количество на страницу
            "page": 0
        }

        vacancies = []
        retries = 3  # Число попыток при ошибках

        while self.running and retries > 0:
            try:
                response = requests.get(base_url, params=params, timeout=10)

                # Обработка ограничения запросов
                if response.status_code == 403:
                    self.update_log.emit("Ошибка: Превышен лимит запросов!")
                    time.sleep(10)
                    continue

                response.raise_for_status()
                data = response.json()

                # Добавление вакансий и проверка страниц
                vacancies.extend(data.get("items", []))
                if params["page"] >= data.get("pages", 1) - 1:
                    break

                params["page"] += 1
                time.sleep(0.25)  # Задержка между запросами
                retries = 3  # Сброс счетчика повторов

            except requests.exceptions.ConnectionError:
                self.update_log.emit("Ошибка подключения к интернету!")
                retries -= 1
                time.sleep(5)

            except requests.exceptions.Timeout:
                self.update_log.emit("Таймаут соединения. Повтор...")
                retries -= 1
                time.sleep(3)

            except Exception as e:
                self.update_log.emit(f"Ошибка запроса: {str(e)}")
                break

        return vacancies

    @staticmethod
    def date_range(start_date: datetime, end_date: datetime,
                   delta: timedelta = timedelta(days=7)) -> Generator[Tuple[datetime, datetime], None, None]:
        """Генератор интервалов дат с шагом delta"""
        current_date = start_date
        while current_date < end_date:
            next_date = min(current_date + delta, end_date)
            yield (current_date, next_date)
            current_date = next_date

    def save_to_excel(self, vacancies: list) -> None:
        """Сохранение результатов в Excel с форматированием"""
        data = []
        for item in vacancies:
            if not item:  # Пропуск пустых элементов
                continue

            # Безопасное получение данных (защита от None)
            salary: Dict[str, Any] = item.get("salary") or {}
            employer: Dict[str, Any] = item.get("employer") or {}
            area: Dict[str, Any] = item.get("area") or {}

            # Парсинг даты публикации
            try:
                pub_date = parser.parse(item.get("published_at", "")).strftime("%d.%m.%Y %H:%M")
            except:
                pub_date = "N/A"

            row = {
                "Название": item.get("name"),
                "Компания": employer.get("name"),
                "Зарплата от": salary.get("from"),
                "Зарплата до": salary.get("to"),
                "Валюта": salary.get("currency"),
                "Регион": area.get("name"),
                "Дата публикации": pub_date,
                "Ссылка": item.get("alternate_url")
            }
            data.append(row)

        # Создание безопасного имени файла
        safe_name = "".join([c if c.isalnum() or c in ('_', '-') else '_' for c in self.profession]).rstrip('_')
        file_name = f"{safe_name}_vacancies.xlsx"

        try:
            # Сохранение через pandas
            df = pd.DataFrame(data)
            with pd.ExcelWriter(file_name, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='Вакансии')

            # Добавление табличного формата
            wb = load_workbook(file_name)
            ws = wb.active
            tab = Table(displayName="VacanciesTable", ref=f"A1:H{len(data) + 1}")
            tab.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium9",
                showRowStripes=True
            )
            ws.add_table(tab)
            wb.save(file_name)

            self.update_log.emit(f"Файл сохранен: {os.path.abspath(file_name)}")

        except Exception as e:
            self.update_log.emit(f"Ошибка сохранения: {str(e)}")

    def stop(self) -> None:
        """Остановка потока"""
        self.running = False
        self.wait(5000)  # Ожидание завершения 5 сек


class MainWindow(QWidget):
    """Главное окно приложения"""

    def __init__(self):
        super().__init__()
        self.parser_thread: Optional[ParserThread] = None
        self.initUI()

    def initUI(self) -> None:
        """Инициализация интерфейса"""
        self.setWindowTitle("Парсер вакансий hh.ru v2.3")
        self.setGeometry(400, 400, 800, 600)

        layout = QVBoxLayout()

        # Элементы интерфейса
        self.progress_bar = QProgressBar()
        self.progress_bar.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.progress_bar.setFormat("Прогресс: %p%")

        self.label = QLabel("Введите профессию для поиска:")
        self.profession_input = QLineEdit()
        self.start_btn = QPushButton("Начать поиск")
        self.log_output = QTextEdit()
        self.log_output.setReadOnly(True)

        # Добавление элементов в layout
        layout.addWidget(self.progress_bar)
        layout.addWidget(self.label)
        layout.addWidget(self.profession_input)
        layout.addWidget(self.start_btn)
        layout.addWidget(self.log_output)

        self.setLayout(layout)
        self.start_btn.clicked.connect(self.start_parsing)

    def start_parsing(self) -> None:
        """Обработчик нажатия кнопки 'Начать поиск'"""
        profession = self.profession_input.text().strip()
        if not profession:
            self.log_output.append("⚠ Введите профессию!")
            return

        if self.parser_thread and self.parser_thread.isRunning():
            self.log_output.append("⚠ Парсинг уже выполняется!")
            return

        # Сброс интерфейса
        self.progress_bar.setValue(0)
        self.log_output.clear()

        # Запуск потока
        self.parser_thread = ParserThread(profession)
        self.parser_thread.update_log.connect(self.log_output.append)
        self.parser_thread.progress_updated.connect(self.progress_bar.setValue)
        self.parser_thread.finished.connect(self.on_finished)
        self.parser_thread.start()

        self.start_btn.setEnabled(False)
        self.start_btn.setText("Идёт поиск...")

    def on_finished(self) -> None:
        """Обработчик завершения работы потока"""
        self.start_btn.setEnabled(True)
        self.start_btn.setText("Начать поиск")
        self.parser_thread = None

    def closeEvent(self, event) -> None:
        """Обработчик закрытия окна"""
        if self.parser_thread and self.parser_thread.isRunning():
            self.parser_thread.stop()
        event.accept()


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())
